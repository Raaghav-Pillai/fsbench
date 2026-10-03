"""Render a Doc into bytes for a given file format."""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone

from fsbench.docs import Doc, Table

FIXED_TIMESTAMP = datetime(2026, 10, 1, 9, 0, 0, tzinfo=timezone.utc)


def table_lines(table: Table, sep: str = " | ") -> list[str]:
    widths = [len(c) for c in table.columns]
    for row in table.rows:
        for i, cell in enumerate(row):
            widths[i] = max(widths[i], len(cell))

    def fmt(cells: list[str]) -> str:
        return sep.join(c.ljust(w) for c, w in zip(cells, widths)).rstrip()

    lines = [fmt(table.columns), fmt(["-" * w for w in widths])]
    lines.extend(fmt(r) for r in table.rows)
    return lines


def render_txt(doc: Doc) -> str:
    lines = [doc.title, "=" * len(doc.title), ""]
    lines.extend(f"{k}: {v}" for k, v in doc.fields)
    if doc.fields:
        lines.append("")
    if doc.table:
        lines.extend(table_lines(doc.table))
        lines.append("")
    for p in doc.paragraphs:
        lines.extend([p, ""])
    return "\n".join(lines).rstrip() + "\n"


def render_md(doc: Doc) -> str:
    lines = [f"# {doc.title}", ""]
    lines.extend(f"- **{k}:** {v}" for k, v in doc.fields)
    if doc.fields:
        lines.append("")
    if doc.table:
        t = doc.table
        lines.append("| " + " | ".join(t.columns) + " |")
        lines.append("|" + "|".join("---" for _ in t.columns) + "|")
        lines.extend("| " + " | ".join(r) + " |" for r in t.rows)
        lines.append("")
    for p in doc.paragraphs:
        lines.extend([p, ""])
    return "\n".join(lines).rstrip() + "\n"


def render_csv(doc: Doc) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    if doc.table:
        buf.write(f"# {doc.title}\n")
        for k, v in doc.fields:
            buf.write(f"# {k}: {v}\n")
        for p in doc.paragraphs:
            buf.write(f"# {p}\n")
        w.writerow(doc.table.columns)
        w.writerows(doc.table.rows)
    else:
        w.writerow(["field", "value"])
        w.writerow(["Title", doc.title])
        w.writerows([k, v] for k, v in doc.fields)
        w.writerows(["Text", p] for p in doc.paragraphs)
    return buf.getvalue()


def render_json(doc: Doc) -> str:
    obj: dict = {"title": doc.title}
    if doc.fields:
        obj["fields"] = {k: v for k, v in doc.fields}
    if doc.table:
        obj["rows"] = [dict(zip(doc.table.columns, r)) for r in doc.table.rows]
    if doc.paragraphs:
        obj["text"] = doc.paragraphs
    return json.dumps(obj, indent=2) + "\n"


def render_pdf(doc: Doc) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_creation_date(FIXED_TIMESTAMP)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 8, _latin1(doc.title), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 10)
    for k, v in doc.fields:
        pdf.multi_cell(0, 5.5, _latin1(f"{k}: {v}"), new_x="LMARGIN", new_y="NEXT")
    if doc.table:
        pdf.ln(3)
        pdf.set_font("Courier", "", 8)
        for line in table_lines(doc.table, sep="  "):
            pdf.multi_cell(0, 4, _latin1(line), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
    for p in doc.paragraphs:
        pdf.ln(2)
        pdf.multi_cell(0, 5.5, _latin1(p), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def render_xlsx(doc: Doc) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.created = FIXED_TIMESTAMP.replace(tzinfo=None)
    wb.properties.modified = FIXED_TIMESTAMP.replace(tzinfo=None)
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([doc.title])
    for k, v in doc.fields:
        ws.append([k, v])
    if doc.table:
        ws.append([])
        ws.append(doc.table.columns)
        for r in doc.table.rows:
            ws.append(r)
    for p in doc.paragraphs:
        ws.append([])
        ws.append([p])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_docx(doc: Doc) -> bytes:
    import docx

    d = docx.Document()
    d.core_properties.created = FIXED_TIMESTAMP.replace(tzinfo=None)
    d.core_properties.modified = FIXED_TIMESTAMP.replace(tzinfo=None)
    d.add_heading(doc.title, level=1)
    for k, v in doc.fields:
        para = d.add_paragraph()
        para.add_run(f"{k}: ").bold = True
        para.add_run(v)
    if doc.table:
        t = d.add_table(rows=1, cols=len(doc.table.columns))
        t.style = "Table Grid"
        for cell, text in zip(t.rows[0].cells, doc.table.columns):
            cell.text = text
        for r in doc.table.rows:
            for cell, text in zip(t.add_row().cells, r):
                cell.text = text
    for p in doc.paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _latin1(text: str) -> str:
    return text.encode("latin-1", "replace").decode("latin-1")


_TEXT_RENDERERS = {"txt": render_txt, "md": render_md, "csv": render_csv, "json": render_json}
_BINARY_RENDERERS = {"pdf": render_pdf, "xlsx": render_xlsx, "docx": render_docx}


def render(doc: Doc, fmt: str) -> bytes:
    if fmt in _TEXT_RENDERERS:
        return _TEXT_RENDERERS[fmt](doc).encode("utf-8")
    if fmt in _BINARY_RENDERERS:
        return _BINARY_RENDERERS[fmt](doc)
    raise ValueError(f"unsupported format: {fmt}")
