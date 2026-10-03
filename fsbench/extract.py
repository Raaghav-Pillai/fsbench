"""Extract readable text from workspace files of any supported format."""

from __future__ import annotations

import io
from pathlib import Path

BINARY_MAGIC = (b"%PDF", b"PK\x03\x04")


def is_binary(data: bytes) -> bool:
    return data.startswith(BINARY_MAGIC) or b"\x00" in data[:4096]


def extract_text(path: Path) -> str:
    data = path.read_bytes()
    suffix = path.suffix.lower()
    if data.startswith(b"%PDF"):
        return _pdf_text(data)
    if data.startswith(b"PK\x03\x04"):
        if suffix == ".xlsx":
            return _xlsx_text(data)
        if suffix == ".docx":
            return _docx_text(data)
        for reader in (_docx_text, _xlsx_text):
            try:
                return reader(data)
            except Exception:
                continue
        raise ValueError("unrecognised zip container")
    if b"\x00" in data[:4096]:
        raise ValueError("binary file with no text extractor")
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip() + "\n"


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        if len(wb.worksheets) > 1:
            lines.append(f"[sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            lines.append("\t".join(cells))
    wb.close()
    return "\n".join(lines).strip() + "\n"


def _docx_text(data: bytes) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    lines = []
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(el, d).text)
        elif tag == "tbl":
            for row in Table(el, d).rows:
                lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines).strip() + "\n"
